import PyPDF2
from docx import Document
import io
import re

class ResumeParser:
    @staticmethod
    async def parse_pdf(file_content: bytes) -> str:
        """Extract text from PDF file using multiple methods"""
        text = ""
        
        # Method 1: Try PyPDF2
        try:
            pdf_reader = PyPDF2.PdfReader(io.BytesIO(file_content))
            for page in pdf_reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
        except Exception as e:
            print(f"PyPDF2 error: {e}")
        
        # Method 2: If no text found, try different approach
        if not text.strip():
            try:
                import pdfplumber
                with pdfplumber.open(io.BytesIO(file_content)) as pdf:
                    for page in pdf.pages:
                        page_text = page.extract_text()
                        if page_text:
                            text += page_text + "\n"
            except ImportError:
                print("pdfplumber not installed")
            except Exception as e:
                print(f"pdfplumber error: {e}")
        
        # Clean up text
        text = ResumeParser.clean_text(text)
        
        if not text.strip():
            raise Exception("Could not extract text from PDF. File might be scanned or image-based.")
        
        return text.strip()
    
    @staticmethod
    async def parse_docx(file_content: bytes) -> str:
        """Extract text from DOCX file"""
        try:
            doc = Document(io.BytesIO(file_content))
            text = ""
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text += paragraph.text + "\n"
            
            # Also extract text from tables
            for table in doc.tables:
                for row in table.rows:
                    for cell in row.cells:
                        if cell.text.strip():
                            text += cell.text + "\n"
            
            return text.strip()
        except Exception as e:
            raise Exception(f"Error parsing DOCX: {str(e)}")
    
    @staticmethod
    def clean_text(text: str) -> str:
        """Clean extracted text"""
        # Remove extra whitespaces
        text = re.sub(r'\s+', ' ', text)
        # Remove special characters but keep important ones
        text = re.sub(r'[^\w\s\.\,\-\#\+\(\)\/\@]', '', text)
        return text.strip()
    
    @staticmethod
    async def extract_text(file_content: bytes, filename: str) -> str:
        """Extract text based on file extension"""
        if filename.lower().endswith('.pdf'):
            return await ResumeParser.parse_pdf(file_content)
        elif filename.lower().endswith('.docx'):
            return await ResumeParser.parse_docx(file_content)
        else:
            raise ValueError(f"Unsupported file type: {filename}")